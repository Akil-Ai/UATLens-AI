import io
import pytest
from openpyxl import load_workbook
from backend.app.export.excel_builder import generate_excel_workbook


def test_excel_export_structure():
    test_cases = [
        {
            "id": "TC-001",
            "scenario": "Verify valid checkout",
            "preconditions": ["User logged in", "Cart has items"],
            "steps": ["1. Go to cart", "2. Click checkout"],
            "test_data": {"amount": 50.00, "item": "Keyboard"},
            "expected_result": "Order placed successfully",
            "scenario_type": "Positive",
            "role": "Registered Customer",
            "priority": "High",
            "requirement_id": "REQ-001",
            "status": "Approved",
            "flags": []
        },
        {
            "id": "TC-002",
            "scenario": "Verify invalid coupon code",
            "preconditions": ["Cart total $10"],
            "steps": ["1. Enter coupon", "2. Click apply"],
            "test_data": {"coupon": "INVALID"},
            "expected_result": "Error shown",
            "scenario_type": "Negative",
            "role": "Guest",
            "priority": "Medium",
            "requirement_id": "REQ-002",
            "status": "Draft",
            "flags": [{"type": "Ambiguous Requirement", "severity": "Medium", "message": "Vague error text"}]
        }
    ]
    requirements = [
        {"id": "REQ-001", "title": "Checkout Workflow"},
        {"id": "REQ-002", "title": "Coupon Processing"}
    ]

    excel_stream = generate_excel_workbook(test_cases, requirements, "Test Project")
    wb = load_workbook(excel_stream)

    # 1. Assert sheet names match required 5-sheet structure
    expected_sheets = ["Test Cases", "Requirements", "Clarifications", "Permission Rules", "Coverage Summary"]
    assert wb.sheetnames == expected_sheets

    # 2. Assert exact column headers on Sheet 1
    ws1 = wb["Test Cases"]
    expected_headers = [
        "ID", "Scenario", "Preconditions", "Steps", "Test Data", "Expected Result",
        "Type", "Role", "Priority", "Requirement ID", "Permission Rules", "Status", "Stale/Blocked", "Flags"
    ]
    actual_headers = [ws1.cell(row=1, column=col).value for col in range(1, len(expected_headers) + 1)]
    assert actual_headers == expected_headers

    # 3. Assert row count
    assert ws1.max_row == 3

    # 4. Assert Sheet 2 Requirements has requirement records
    ws2 = wb["Requirements"]
    assert ws2.cell(row=1, column=1).value == "Requirement ID"
    assert ws2.cell(row=2, column=1).value == "REQ-001"

    # 5. Assert Sheet 3 Clarifications has flags / clarifications
    ws3 = wb["Clarifications"]
    assert ws3.max_row >= 2
    assert ws3.cell(row=2, column=1).value == "TC-002"

    # 6. Assert Sheet 5 Coverage Summary has metric rows
    ws5 = wb["Coverage Summary"]
    assert ws5.cell(row=1, column=1).value == "Metric Category"


def test_csv_formula_injection_and_zip():
    from backend.app.export.csv_builder import (
        sanitize_formula_injection,
        generate_standard_csv,
        generate_zip_bundle
    )
    import zipfile

    # Formula injection strings
    dangerous = "=cmd|' /C calc'!A0"
    safe = sanitize_formula_injection(dangerous)
    assert safe.startswith("'")
    assert safe == "'=cmd|' /C calc'!A0"

    plus_formula = "+12345"
    assert sanitize_formula_injection(plus_formula).startswith("'")

    normal = "Normal Text"
    assert sanitize_formula_injection(normal) == "Normal Text"

    test_cases = [{
        "id": "TC-001",
        "scenario": "=HYPERLINK(\"http://evil.com\")",
        "preconditions": [],
        "steps": ["Step 1"],
        "test_data": {},
        "expected_result": "Result",
        "scenario_type": "Positive",
        "role": "User",
        "priority": "High",
        "requirement_id": "REQ-001",
        "permission_rule_ids": ["PR-001"],
        "status": "Draft",
        "flags": []
    }]

    csv_data = generate_standard_csv(test_cases).getvalue()
    assert "'=HYPERLINK" in csv_data

    zip_buf = generate_zip_bundle(test_cases, [{"id": "REQ-001", "title": "Req"}], [], [], "TestProject")
    with zipfile.ZipFile(zip_buf, "r") as zf:
        namelist = zf.namelist()
        assert any("test_cases.csv" in name for name in namelist)
        assert any("requirements.csv" in name for name in namelist)

