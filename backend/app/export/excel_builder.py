import io
import json
from typing import List, Dict, Any, Optional
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DANGEROUS_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def sanitize_cell(val: Any) -> Any:
    """Neutralizes formula injection in Excel cells."""
    if val is None:
        return ""
    if isinstance(val, (int, float, bool)):
        return val
    s = str(val)
    if s.startswith(DANGEROUS_PREFIXES) and not (s.startswith("-") and s[1:].replace(".", "", 1).isdigit()):
        return f"'{s}"
    return s


def generate_excel_workbook(
    test_cases: List[Dict[str, Any]],
    requirements: List[Dict[str, Any]] = None,
    project_name: str = "UATlens Test Suite",
    clarifications: List[Dict[str, Any]] = None,
    permission_rules: List[Dict[str, Any]] = None,
    coverage_summary: Dict[str, Any] = None,
) -> io.BytesIO:
    """
    Generates a professionally styled 5-sheet Excel (.xlsx) workbook with formula injection protection:
    Sheet 1: Test Cases
    Sheet 2: Requirements
    Sheet 3: Clarifications
    Sheet 4: Permission Rules
    Sheet 5: Coverage Summary
    """
    # Backwards compatibility check if caller passes clarifications as 3rd arg
    if isinstance(project_name, list) and clarifications is None:
        clarifications = project_name
        project_name = "UATlens Test Suite"

    if isinstance(clarifications, str) and (project_name == "UATlens Test Suite" or not project_name):
        project_name = clarifications
        clarifications = []

    if requirements is None:
        requirements = []
    if not isinstance(clarifications, list):
        clarifications = []
    if not clarifications:
        for tc in test_cases:
            for f in tc.get("flags", []):
                clarifications.append({
                    "id": tc.get("id"),
                    "requirement_id": tc.get("requirement_id"),
                    "issue_type": f.get("type", "Flag") if isinstance(f, dict) else str(f),
                    "suggested_question": f.get("suggested_question", f.get("message", "")) if isinstance(f, dict) else "",
                    "status": "Open",
                    "reviewer_answer": "",
                    "confirmed_by": "",
                    "dismissal_reason": ""
                })
    if not isinstance(permission_rules, list):
        permission_rules = []
    if coverage_summary is None:
        coverage_summary = {}

    wb = Workbook()

    indigo_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10, color="1E293B")
    center_align = Alignment(horizontal="center", vertical="top", wrap_text=True)
    left_wrap_align = Alignment(horizontal="left", vertical="top", wrap_text=True)

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    # ──────────────── SHEET 1: TEST CASES ────────────────
    ws1 = wb.active
    ws1.title = "Test Cases"
    ws1.sheet_view.showGridLines = True

    tc_headers = [
        "ID", "Scenario", "Preconditions", "Steps", "Test Data", "Expected Result",
        "Type", "Role", "Priority", "Requirement ID", "Permission Rules", "Status", "Stale/Blocked", "Flags"
    ]
    ws1.append(tc_headers)

    for col_idx in range(1, len(tc_headers) + 1):
        cell = ws1.cell(row=1, column=col_idx)
        cell.fill = indigo_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
    ws1.row_dimensions[1].height = 28
    ws1.freeze_panes = "A2"
    if ws1.sheet_view and ws1.sheet_view.selection:
        ws1.sheet_view.selection[0].activeCell = "A2"
        ws1.sheet_view.selection[0].sqref = "A2"

    for row_idx, tc in enumerate(test_cases, start=2):
        preconds = tc.get("preconditions", [])
        preconds_str = "\n".join(preconds) if isinstance(preconds, list) else str(preconds)

        steps = tc.get("steps", [])
        if isinstance(steps, list):
            steps_str = "\n".join([s if s.strip().startswith(tuple(str(k) for k in range(1, 10))) else f"{i+1}. {s}" for i, s in enumerate(steps)])
        else:
            steps_str = str(steps)

        test_data = tc.get("test_data", {})
        test_data_str = "\n".join([f"{k}: {v}" for k, v in test_data.items()]) if isinstance(test_data, dict) else str(test_data)

        perm_rules_str = ", ".join(tc.get("permission_rule_ids", [])) if isinstance(tc.get("permission_rule_ids"), list) else ""
        flags = tc.get("flags", [])
        flags_str = ", ".join([f.get("type", "") for f in flags if isinstance(f, dict)])

        stale_blocked = []
        if tc.get("is_stale"):
            stale_blocked.append("STALE")
        if tc.get("is_blocked"):
            stale_blocked.append(f"BLOCKED: {tc.get('blocked_reason') or 'Unresolved rule'}")
        status_note = " | ".join(stale_blocked) if stale_blocked else "Normal"

        row_data = [
            sanitize_cell(tc.get("id", f"TC-{row_idx-1:03d}")),
            sanitize_cell(tc.get("scenario", "")),
            sanitize_cell(preconds_str),
            sanitize_cell(steps_str),
            sanitize_cell(test_data_str),
            sanitize_cell(tc.get("expected_result", "")),
            sanitize_cell(tc.get("scenario_type", "Positive")),
            sanitize_cell(tc.get("role", "Guest")),
            sanitize_cell(tc.get("priority", "Medium")),
            sanitize_cell(tc.get("requirement_id", "REQ-001")),
            sanitize_cell(perm_rules_str),
            sanitize_cell(tc.get("status", "Draft")),
            sanitize_cell(status_note),
            sanitize_cell(flags_str),
        ]
        ws1.append(row_data)
        for col_idx in range(1, len(row_data) + 1):
            c = ws1.cell(row=row_idx, column=col_idx)
            c.font = data_font
            c.border = thin_border
            c.alignment = center_align if col_idx in (1, 7, 8, 9, 10, 12) else left_wrap_align

    # ──────────────── SHEET 2: REQUIREMENTS ────────────────
    ws2 = wb.create_sheet(title="Requirements")
    ws2.sheet_view.showGridLines = True
    req_headers = ["Requirement ID", "Title", "Text", "Source Quote", "Roles Involved", "Expected Outcome"]
    ws2.append(req_headers)
    for col_idx in range(1, len(req_headers) + 1):
        cell = ws2.cell(row=1, column=col_idx)
        cell.fill = indigo_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
    ws2.row_dimensions[1].height = 28
    ws2.freeze_panes = "A2"
    if ws2.sheet_view and ws2.sheet_view.selection:
        ws2.sheet_view.selection[0].activeCell = "A2"
        ws2.sheet_view.selection[0].sqref = "A2"

    for r_idx, r in enumerate(requirements, start=2):
        roles_str = ", ".join(r.get("roles_involved", [])) if isinstance(r.get("roles_involved"), list) else str(r.get("roles_involved", ""))
        r_data = [
            sanitize_cell(r.get("id")),
            sanitize_cell(r.get("title")),
            sanitize_cell(r.get("text")),
            sanitize_cell(r.get("source_quote")),
            sanitize_cell(roles_str),
            sanitize_cell(r.get("expected_outcome")),
        ]
        ws2.append(r_data)
        for col_idx in range(1, len(r_data) + 1):
            c = ws2.cell(row=r_idx, column=col_idx)
            c.font = data_font
            c.border = thin_border
            c.alignment = center_align if col_idx == 1 else left_wrap_align

    # ──────────────── SHEET 3: CLARIFICATIONS ────────────────
    ws3 = wb.create_sheet(title="Clarifications")
    ws3.sheet_view.showGridLines = True
    clar_headers = ["ID", "Requirement ID", "Issue Type", "Question", "Status", "Authoritative Answer", "Confirmed By", "Dismissal Reason"]
    ws3.append(clar_headers)
    for col_idx in range(1, len(clar_headers) + 1):
        cell = ws3.cell(row=1, column=col_idx)
        cell.fill = indigo_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
    ws3.row_dimensions[1].height = 28
    ws3.freeze_panes = "A2"
    if ws3.sheet_view and ws3.sheet_view.selection:
        ws3.sheet_view.selection[0].activeCell = "A2"
        ws3.sheet_view.selection[0].sqref = "A2"

    for c_idx, c in enumerate(clarifications, start=2):
        c_data = [
            sanitize_cell(c.get("id")),
            sanitize_cell(c.get("requirement_id")),
            sanitize_cell(c.get("issue_type")),
            sanitize_cell(c.get("suggested_question")),
            sanitize_cell(c.get("status")),
            sanitize_cell(c.get("reviewer_answer")),
            sanitize_cell(c.get("confirmed_by")),
            sanitize_cell(c.get("dismissal_reason")),
        ]
        ws3.append(c_data)
        for col_idx in range(1, len(c_data) + 1):
            cell = ws3.cell(row=c_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = center_align if col_idx in (1, 2, 5) else left_wrap_align

    # ──────────────── SHEET 4: PERMISSION RULES ────────────────
    ws4 = wb.create_sheet(title="Permission Rules")
    ws4.sheet_view.showGridLines = True
    perm_headers = ["ID", "Role", "Action", "Resource", "Scope", "Condition", "Decision", "Review Status", "Revision"]
    ws4.append(perm_headers)
    for col_idx in range(1, len(perm_headers) + 1):
        cell = ws4.cell(row=1, column=col_idx)
        cell.fill = indigo_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
    ws4.row_dimensions[1].height = 28
    ws4.freeze_panes = "A2"
    if ws4.sheet_view and ws4.sheet_view.selection:
        ws4.sheet_view.selection[0].activeCell = "A2"
        ws4.sheet_view.selection[0].sqref = "A2"

    for p_idx, p in enumerate(permission_rules, start=2):
        p_data = [
            sanitize_cell(p.get("id")),
            sanitize_cell(p.get("role")),
            sanitize_cell(p.get("action")),
            sanitize_cell(p.get("resource")),
            sanitize_cell(p.get("scope")),
            sanitize_cell(p.get("condition")),
            sanitize_cell(p.get("decision")),
            sanitize_cell(p.get("review_status")),
            sanitize_cell(p.get("revision")),
        ]
        ws4.append(p_data)
        for col_idx in range(1, len(p_data) + 1):
            cell = ws4.cell(row=p_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = center_align if col_idx in (1, 2, 7, 8, 9) else left_wrap_align

    # ──────────────── SHEET 5: COVERAGE SUMMARY ────────────────
    ws5 = wb.create_sheet(title="Coverage Summary")
    ws5.sheet_view.showGridLines = True
    cov_headers = ["Metric Category", "Metric Name", "Value", "Scope / Details"]
    ws5.append(cov_headers)
    for col_idx in range(1, len(cov_headers) + 1):
        cell = ws5.cell(row=1, column=col_idx)
        cell.fill = indigo_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
    ws5.row_dimensions[1].height = 28

    cov_rows = [
        ["Project", "Project Name", project_name, "Exported project"],
        ["Scope", "Export Scope", coverage_summary.get("scope", "Full Project"), "Filter constraints"],
        ["Requirements", "Total Requirements", str(len(requirements)), "All project requirements"],
        ["Requirements", "Tested Requirements", str(coverage_summary.get("tested_requirements", 0)), "Requirements with >=1 test"],
        ["Test Cases", "Total Cases Exported", str(len(test_cases)), "Matching export filters"],
        ["Test Cases", "Approved Cases", str(sum(1 for tc in test_cases if tc.get("status") == "Approved")), "Production-ready"],
        ["Permissions", "Total Active Rules", str(len(permission_rules)), "Active role rules"],
        ["Permissions", "Tested Permissions", str(coverage_summary.get("tested_permissions", 0)), "Rules verified in suite"],
        ["Clarifications", "Resolved Clarifications", str(sum(1 for c in clarifications if c.get("status") == "Resolved")), "Confirmed by reviewer"],
    ]

    for c_row_idx, row_items in enumerate(cov_rows, start=2):
        ws5.append([sanitize_cell(x) for x in row_items])
        for col_idx in range(1, len(row_items) + 1):
            cell = ws5.cell(row=c_row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = center_align if col_idx in (1, 3) else left_wrap_align

    # Auto-adjust column widths
    for sheet in [ws1, ws2, ws3, ws4, ws5]:
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col[:40]:  # check first 40 rows
                if cell.value:
                    val_str = str(cell.value)
                    first_line = val_str.split("\n")[0]
                    max_len = max(max_len, len(first_line))
            sheet.column_dimensions[col_letter].width = max(12, min(max_len + 4, 45))

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
